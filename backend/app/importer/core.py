"""Lógica del importador de `Inventario-80` (ver docs/DATOS.md).

`import_inventory` es la función principal: lee el Excel fila por fila, aplica las
reglas de limpieza de `cleaning.py`, resuelve conflictos de posición y persiste
usuarios / cajas / muestras / movimientos. Nunca aborta por una fila mala: la
reporta como anomalía y sigue con la siguiente.
"""

from __future__ import annotations

import hashlib
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.importer.cleaning import (
    clean_environ_id,
    clean_text,
    clean_tipo,
    parse_caja_numero,
    parse_encargado,
    parse_fecha_entrada,
    parse_fecha_salida,
    parse_pasaje,
    parse_posicion,
    parse_propietario_caja,
    parse_rack_letter,
    parse_seccion,
    parse_si_no,
)
from app.importer.report import Anomaly
from app.models import (
    Box,
    ImportAnomaly,
    ImportRun,
    Movement,
    MovementAction,
    Rack,
    Sample,
    SampleStatus,
    User,
)

SHEET_NAME = "Inventario-80"
HEADER_ROW = 2
DATA_START_ROW = 3

COLUMNS = [
    "ID Environ",
    "ID Origen o Descripción",
    "Caja origen",
    "Tipo",
    "Encargado",
    "Pasaje",
    "Nucleo",
    "Seccion",
    "Rack",
    "Caja",
    "Posición",
    "Fecha de entrada",
    "Caja Completa Si/No",
    "Propietario de Caja",
    "Fecha de salida",
    "Comentarios",
]


@dataclass
class ImportSummary:
    total_rows: int = 0
    imported: int = 0
    withdrawn: int = 0
    skipped_already_imported: int = 0
    skipped_invalid: int = 0
    conflicts_resolved: int = 0
    anomalies: int = 0


@dataclass
class ImportResult:
    summary: ImportSummary
    anomalies: list[Anomaly]
    #: Id de la corrida registrada. `None` en `--dry-run`, que no escribe nada.
    run_id: int | None = None


@dataclass
class _RowRecord:
    row_num: int
    environ_id: str | None
    description: str | None
    sample_type: str
    type_other: str | None
    owner_initials: list[str]
    passage: int | None
    is_core: bool | None
    rack_letter: str
    box_number: int
    position: str
    box_type: str
    box_label: str | None
    box_is_full: bool | None
    box_owner_initials: str | None
    entry_date: date | None
    exit_date: date | None
    exit_initials: str | None
    exit_note: str | None
    notes: str | None
    status: str
    box_id: int = 0


def _raw_str(value: object) -> str:
    return "" if value is None else str(value)


def _find_oversized_field(
    fields: list[tuple[str, str | None, int]]
) -> tuple[str, str, int] | None:
    for column, value, max_len in fields:
        if value is not None and len(value) > max_len:
            return column, value, max_len
    return None


def _source_file_identity(path: Path) -> str:
    """Identidad estable del archivo de origen: hash del contenido, no el nombre.

    Dos workbooks distintos con el mismo nombre no deben compartir idempotencia,
    y el mismo workbook renombrado sigue siendo el mismo origen (ver revisión de
    PR #14: `Path(path).name` colisionaba entre workbooks distintos)."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def import_inventory(path: str | Path, session: Session, *, dry_run: bool = False) -> ImportResult:
    path = Path(path)
    source_file = _source_file_identity(path)

    # En modo read_only openpyxl mantiene el archivo abierto hasta close(): se leen las
    # filas a memoria y se cierra, o en Windows el Excel queda bloqueado.
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"La hoja '{SHEET_NAME}' no existe en {path}")
        sheet = workbook[SHEET_NAME]
        header_values = next(sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))
        data_rows = list(sheet.iter_rows(min_row=DATA_START_ROW, values_only=True))
    finally:
        workbook.close()

    column_index = {
        str(value).strip(): idx for idx, value in enumerate(header_values) if value is not None
    }
    missing = [column for column in COLUMNS if column not in column_index]
    if missing:
        raise ValueError(f"Faltan columnas esperadas en la hoja: {', '.join(missing)}")

    anomalies: list[Anomaly] = []
    summary = ImportSummary()

    racks_by_letter = {rack.letter: rack for rack in session.query(Rack).all()}
    users_by_initials = {user.initials: user for user in session.query(User).all()}
    boxes_by_key: dict[tuple[int, int], Box] = {
        (box.rack_id, box.number): box for box in session.query(Box).all()
    }
    existing_source_rows = {
        row[0]
        for row in session.query(Sample.source_row).filter(
            Sample.source_file == source_file, Sample.source_row.isnot(None)
        )
    }

    def get_or_create_user(initials: str) -> User:
        user = users_by_initials.get(initials)
        if user is None:
            user = User(initials=initials)
            session.add(user)
            session.flush()
            users_by_initials[initials] = user
        return user

    records: list[_RowRecord] = []
    for row_num, row_values in enumerate(data_rows, start=DATA_START_ROW):
        values = {
            name: row_values[idx] if idx < len(row_values) else None for name, idx in column_index.items()
        }
        if all(value is None for value in values.values()):
            continue

        summary.total_rows += 1

        # La fila se parsea SIEMPRE, incluso si ya se importó: las anomalías son una
        # propiedad del Excel, no del estado de la base. Saltarse el parseo hacía que
        # reimportar el mismo archivo reportara ~cero anomalías, y la serie de calidad de
        # datos mostrara una caída a cero que se lee como "lo arreglamos" con el dato
        # igual de sucio. Lo que se saltea es la PERSISTENCIA, que es lo que garantiza la
        # idempotencia que promete el README.
        row_anomalies: list[Anomaly] = []
        record = _parse_row(row_num, values, row_anomalies, racks_by_letter)
        anomalies.extend(row_anomalies)

        if row_num in existing_source_rows:
            summary.skipped_already_imported += 1
            continue
        if record is None:
            summary.skipped_invalid += 1
            continue
        records.append(record)

    _resolve_new_row_conflicts(records, anomalies, summary)

    # Una fila cuya posición no corresponde al tipo de su caja (p. ej. "15" en una caja de
    # cartón ya registrada desde la página) no se importa: quedaría una muestra en una
    # posición inexistente, contada en la ocupación pero invisible en el visor.
    placed_records: list[_RowRecord] = []
    for record in records:
        rack = racks_by_letter[record.rack_letter]
        box_key = (rack.id, record.box_number)
        box = boxes_by_key.get(box_key)
        if box is None:
            box = Box(
                rack_id=rack.id,
                number=record.box_number,
                box_type=record.box_type,
                label=record.box_label,
                is_full=record.box_is_full,
            )
            if record.box_owner_initials:
                box.owner = get_or_create_user(record.box_owner_initials)
            session.add(box)
            session.flush()
            boxes_by_key[box_key] = box
        elif box.box_type != record.box_type:
            anomalies.append(
                Anomaly(
                    row=record.row_num,
                    column="Posición",
                    value=record.position,
                    reason=(
                        f"Tipo de caja inconsistente: la caja {rack.letter}{record.box_number} "
                        f"ya es '{box.box_type}'; la fila no se importa"
                    ),
                )
            )
            summary.skipped_invalid += 1
            continue
        record.box_id = box.id
        placed_records.append(record)
    records = placed_records

    _resolve_conflicts_with_existing_active(session, records, anomalies, summary)

    for record in records:
        sample = Sample(
            environ_id=record.environ_id,
            description=record.description,
            type=record.sample_type,
            type_other=record.type_other,
            owners=[get_or_create_user(initials) for initials in record.owner_initials],
            passage=record.passage,
            is_core=record.is_core,
            status=record.status,
            box_id=record.box_id,
            position=record.position,
            notes=record.notes,
            source_file=source_file,
            source_row=record.row_num,
        )
        session.add(sample)
        session.flush()

        movement_date = record.entry_date or record.exit_date or date.today()
        session.add(
            Movement(
                sample_id=sample.id,
                action=MovementAction.FREEZE.value,
                date=movement_date,
                box_id=record.box_id,
                position=record.position,
            )
        )

        if record.status == SampleStatus.WITHDRAWN.value:
            operator = get_or_create_user(record.exit_initials) if record.exit_initials else None
            session.add(
                Movement(
                    sample_id=sample.id,
                    action=MovementAction.THAW.value,
                    date=record.exit_date or movement_date,
                    operator_id=operator.id if operator else None,
                    box_id=record.box_id,
                    position=record.position,
                    note=record.exit_note,
                )
            )
            summary.withdrawn += 1

        summary.imported += 1

    summary.anomalies = len(anomalies)

    if dry_run:
        # El flag promete no escribir nada, y eso incluye la corrida: un ensayo no es un
        # hecho del historial.
        session.rollback()
        return ImportResult(summary=summary, anomalies=anomalies)

    # La corrida y sus anomalías van en la MISMA transacción que las muestras. Si se
    # commitearan aparte y fallara la segunda, quedaría un inventario importado sin
    # registro de qué pasó al importarlo.
    run = ImportRun(
        source_file=source_file,
        source_name=path.name,
        dry_run=False,
        total_rows=summary.total_rows,
        imported=summary.imported,
        withdrawn=summary.withdrawn,
        skipped_already_imported=summary.skipped_already_imported,
        skipped_invalid=summary.skipped_invalid,
        conflicts_resolved=summary.conflicts_resolved,
        anomalies_count=summary.anomalies,
    )
    session.add(run)
    session.flush()
    for anomaly in anomalies:
        session.add(
            ImportAnomaly(
                run_id=run.id,
                row=anomaly.row,
                column=anomaly.column,
                value=anomaly.value,
                reason=anomaly.reason,
            )
        )
    session.commit()

    return ImportResult(summary=summary, anomalies=anomalies, run_id=run.id)


def _parse_row(
    row_num: int,
    values: dict[str, object],
    anomalies: list[Anomaly],
    racks_by_letter: dict[str, Rack],
) -> _RowRecord | None:
    environ_id, reason = clean_environ_id(values.get("ID Environ"))
    if reason:
        anomalies.append(Anomaly(row_num, "ID Environ", _raw_str(values.get("ID Environ")), reason))

    sample_type, type_other, reason = clean_tipo(values.get("Tipo"))
    if reason:
        anomalies.append(Anomaly(row_num, "Tipo", _raw_str(values.get("Tipo")), reason))
    if sample_type is None:
        return None

    owner_initials = parse_encargado(values.get("Encargado"))

    is_core = parse_si_no(values.get("Nucleo"))
    if is_core is None:
        anomalies.append(Anomaly(row_num, "Nucleo", _raw_str(values.get("Nucleo")), "Núcleo desconocido"))

    seccion, seccion_reason = parse_seccion(values.get("Seccion"))

    rack_letter, rack_reason = parse_rack_letter(values.get("Rack"))
    if rack_reason:
        anomalies.append(Anomaly(row_num, "Rack", _raw_str(values.get("Rack")), rack_reason))
        return None

    rack = racks_by_letter.get(rack_letter)
    if rack is None:
        anomalies.append(
            Anomaly(
                row_num,
                "Rack",
                rack_letter,
                "Rack no configurado en backend/app/seed/layout.yaml; no se crea automáticamente",
            )
        )
        return None

    if seccion_reason:
        anomalies.append(Anomaly(row_num, "Seccion", _raw_str(values.get("Seccion")), seccion_reason))
    elif seccion != rack.section.code:
        anomalies.append(
            Anomaly(
                row_num,
                "Seccion",
                _raw_str(values.get("Seccion")),
                f"No coincide con la sección real del rack {rack_letter} ({rack.section.code})",
            )
        )

    box_number, box_reason = parse_caja_numero(values.get("Caja"))
    if box_reason:
        anomalies.append(Anomaly(row_num, "Caja", _raw_str(values.get("Caja")), box_reason))
        return None

    position, box_type, position_reason = parse_posicion(values.get("Posición"))
    if position_reason:
        anomalies.append(Anomaly(row_num, "Posición", _raw_str(values.get("Posición")), position_reason))
        return None

    entry_date, entry_reason = parse_fecha_entrada(values.get("Fecha de entrada"))
    if entry_reason:
        anomalies.append(
            Anomaly(row_num, "Fecha de entrada", _raw_str(values.get("Fecha de entrada")), entry_reason)
        )

    passage, passage_reason = parse_pasaje(values.get("Pasaje"))
    if passage_reason:
        anomalies.append(Anomaly(row_num, "Pasaje", _raw_str(values.get("Pasaje")), passage_reason))

    exit_date, exit_initials, exit_note, exit_reason = parse_fecha_salida(values.get("Fecha de salida"))
    if exit_reason:
        anomalies.append(
            Anomaly(row_num, "Fecha de salida", _raw_str(values.get("Fecha de salida")), exit_reason)
        )

    status = (
        SampleStatus.WITHDRAWN.value
        if clean_text(values.get("Fecha de salida")) is not None
        else SampleStatus.ACTIVE.value
    )

    description = clean_text(values.get("ID Origen o Descripción"))
    box_label = clean_text(values.get("Caja origen"))
    box_owner_initials = parse_propietario_caja(values.get("Propietario de Caja"))

    # Los límites siguen a las columnas de los modelos (Sample.environ_id/description/
    # type_other, Box.label, User.initials): sin esta validación una fila con un valor
    # demasiado largo aborta el flush completo en vez de reportarse y saltarse.
    oversized = _find_oversized_field(
        [
            ("ID Environ", environ_id, 60),
            ("ID Origen o Descripción", description, 255),
            ("Tipo", type_other, 120),
            ("Caja origen", box_label, 120),
            *(("Encargado", initials, 10) for initials in owner_initials),
            ("Propietario de Caja", box_owner_initials, 10),
            ("Fecha de salida", exit_initials, 10),
        ]
    )
    if oversized:
        column, value, max_len = oversized
        anomalies.append(
            Anomaly(row_num, column, value, f"Valor de más de {max_len} caracteres, no se pudo importar")
        )
        return None

    return _RowRecord(
        row_num=row_num,
        environ_id=environ_id,
        description=description,
        sample_type=sample_type,
        type_other=type_other,
        owner_initials=owner_initials,
        passage=passage,
        is_core=is_core,
        rack_letter=rack_letter,
        box_number=box_number,
        position=position,
        box_type=box_type,
        box_label=box_label,
        box_is_full=parse_si_no(values.get("Caja Completa Si/No")),
        box_owner_initials=box_owner_initials,
        entry_date=entry_date,
        exit_date=exit_date,
        exit_initials=exit_initials,
        exit_note=exit_note,
        notes=clean_text(values.get("Comentarios")),
        status=status,
    )


def _resolve_new_row_conflicts(
    records: list[_RowRecord], anomalies: list[Anomaly], summary: ImportSummary
) -> None:
    """Entre las filas nuevas de esta corrida: si más de una queda activa en la
    misma posición, gana la fila más reciente (número de fila más alto) y las
    demás se importan como retiradas por conflicto (ver docs/DATOS.md)."""
    groups: dict[tuple[str, int, str], list[_RowRecord]] = defaultdict(list)
    for record in records:
        if record.status == SampleStatus.ACTIVE.value:
            groups[(record.rack_letter, record.box_number, record.position)].append(record)

    for group in groups.values():
        if len(group) <= 1:
            continue
        group.sort(key=lambda item: item.row_num)
        winner = group[-1]
        for loser in group[:-1]:
            loser.status = SampleStatus.WITHDRAWN.value
            loser.exit_date = loser.exit_date or loser.entry_date
            loser.exit_initials = None
            loser.exit_note = f"Retirada automáticamente: conflicto de posición con la fila {winner.row_num}"
            anomalies.append(
                Anomaly(
                    row=loser.row_num,
                    column="Posición",
                    value=loser.position,
                    reason=f"Conflicto de posición con la fila {winner.row_num}: se retiró automáticamente",
                )
            )
            summary.conflicts_resolved += 1


def _resolve_conflicts_with_existing_active(
    session: Session, records: list[_RowRecord], anomalies: list[Anomaly], summary: ImportSummary
) -> None:
    """Una fila nueva no puede quedar activa en una posición ya ocupada por una
    muestra activa de una importación anterior: se retira automáticamente."""
    active_records = [record for record in records if record.status == SampleStatus.ACTIVE.value]
    if not active_records:
        return

    box_ids = {record.box_id for record in active_records}
    existing_active = (
        session.query(Sample)
        .filter(Sample.status == SampleStatus.ACTIVE.value, Sample.box_id.in_(box_ids))
        .all()
    )
    existing_by_key = {(sample.box_id, sample.position): sample for sample in existing_active}

    for record in active_records:
        existing = existing_by_key.get((record.box_id, record.position))
        if existing is None:
            continue
        record.status = SampleStatus.WITHDRAWN.value
        record.exit_date = record.exit_date or record.entry_date
        record.exit_initials = None
        record.exit_note = (
            f"Retirada automáticamente: conflicto de posición con la muestra activa "
            f"id={existing.id} (fila {existing.source_row})"
        )
        anomalies.append(
            Anomaly(
                row=record.row_num,
                column="Posición",
                value=record.position,
                reason=(
                    f"Conflicto de posición con una muestra ya activa (id={existing.id}, "
                    f"fila {existing.source_row}): se retiró automáticamente"
                ),
            )
        )
        summary.conflicts_resolved += 1
